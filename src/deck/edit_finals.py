from pathlib import Path as _P
DECK_ROOT = str(_P(__file__).resolve().parents[2])
exec(open(_P(__file__).with_name("helpers_block.py"), encoding="utf-8").read())
from pptx.util import Pt
INK = "4A4D55"
# S[n] is now the slide at position n. Main flow 1-10, appendix 11-17.

# ---------- numbering: eyebrow (shape 12) and page number ("n / 15") ----------
SECTION = {2:"P R O B L E M",3:"S O L U T I O N",4:"V A L I D A T I O N",5:"P R O O F   I N   R U P E E S",6:"W O R K I N G   D E M O",
           7:"L A N D E D   C O S T",8:"P I L O T   P L A N",9:"R I S K S   &   R O A D M A P"}
APPX = {11:"D R I V E R S   &   D A T A",12:"D A T A   &   L E A K A G E   C O N T R O L S",13:"D R I V E R   S T U D Y",
        14:"M O D E L S",15:"E X P L A I N A B I L I T Y",16:"V A L I D A T I O N   M E T H O D",17:"U N C E R T A I N T Y   B A N D S"}
for n, sl in S.items():
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text.strip()
        if re.fullmatch(r"\d+ / 15", t) and sh.top > Inches(6.8):
            T(n, sh.shape_id, f"{n} / 10" if n <= 10 else f"A{n-10}")
        if sh.shape_id == 12 and "·" in t and n not in (1, 10):
            if n <= 10: T(n, 12, f"0 {n}  ·  {SECTION[n]}" if n < 10 else f"1 0  ·  {SECTION[n]}")
            else: T(n, 12, f"A {n-10}  ·  A P P E N D I X  ·  {APPX[n]}")

def move_pill(n, to_label):
    sl = S[n]; pills = {sh.text_frame.text.strip(): sh for sh in sl.shapes if sh.has_text_frame and sh.top < Inches(0.6) and sh.width < Inches(1.9)}
    red = [sh for sh in sl.shapes if sh.top < Inches(0.5) and sh.shape_type == 1 and sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == "C8102E"][0]
    old = min(pills.values(), key=lambda p: abs(p.left - red.left))
    other_col = pills["Problem"].text_frame._txBody.find(".//" + qn("a:srgbClr")).get("val")
    red.left = pills[to_label].left + (red.left - old.left)
    for sh in sl.shapes:
        if sh.has_text_frame and sh.top < Inches(0.6) and sh.text_frame.text.strip() not in ("", to_label):
            _fill_body(sh.text_frame._txBody, [sh.text_frame.text], other_col)
    _fill_body(pills[to_label].text_frame._txBody, [to_label], "FFFFFF")

def notes(n, text):
    S[n].notes_slide.notes_text_frame.text = text

# ============================ 1. OPENING (problem + impact in 60 s) ============================
T(1, 11, "T H E   B U S I N E S S   P R O B L E M")
T(1, 12, "Aluminium foil and PVC resin are about 70% of the cost of a pharma blister pack, and both are bought "
         "on judgement with no forward view of price. Aluminium swung 2.7× between 2020 and 2022, and a 1% move in "
         "raw material price moves total manufacturing cost by about 0.7%.")
T(1, 14, "₹0.45 Cr")
T(1, 15, "saved per ₹100 Cr of PVC spend in", "a 78-month backtest, net of carry")
T(1, 17, "+20.7%")
T(1, 18, "PVC forecast skill vs a no-change", "baseline at 1 month (p = 0.0004)")
T(1, 20, "₹0 capex")
T(1, 21, "90-day shadow pilot: buyers keep", "deciding while we measure")
T(1, 9, "Finals  ·  Technology Category")
cell(1, 27, 4, 1, "Live dashboard, 14-slide deck + appendix, reproducible pipeline with leakage tests")
T(1, 23, "W H A T   W E   B U I L T   A N D   P R O V E D")
T(1, 24, "Smart Buy forecasts both prices six months ahead and tells buyers when to buy ahead, buy normal or defer, "
         "net of carrying cost. Back-tested month by month from 2018 to 2024 under a leak-proof test, PVC timing beats "
         "random (p = 0.002). Aluminium is near-efficient, so there the value is budgeting and risk bands, not timing. "
         "We are asking for a zero-capex shadow pilot.")
notes(1, "OPENING, 60 SECONDS.\n"
      "[Problem, 15 s] Seventy percent of what a blister pack costs ACG is aluminium foil and PVC resin. Both are bought on judgement. "
      "Aluminium swung 2.7 times in two years; every 1% move in these prices moves total manufacturing cost by about 0.7%.\n"
      "[What we built, 20 s] Smart Buy forecasts both prices six months ahead and turns that into one instruction for the buyer: "
      "buy ahead, buy normal, or defer, after the cost of holding stock.\n"
      "[Proof, 15 s] We replayed it on every month from 2018 to 2024. On PVC it would have saved about 0.45 Cr per 100 Cr of spend, "
      "and the timing beats random with p = 0.002. On aluminium we tell you honestly: the market is near-efficient, so we use it for budgets and risk.\n"
      "[Ask, 10 s] We are asking for a 90-day, zero-capex shadow pilot. Buyers keep deciding; we measure.")

# ============================ 7. LANDED COST (new, from the data-lag template) ============================
move_pill(7, "Implementation")
T(7, 13, "Landed cost = global index + five gaps. We forecast the index; the gaps",
         "are rules, contracts or dated events that we calculate")
T(7, 14, "From the global index to ACG's landed price")
tbl = shp(7, 15).table
for j, w in enumerate((1.75, 3.25, 2.30)):
    tbl.columns[j].width = Inches(w)
rows = [("Component", "What sets it, and how often it moves", "Size today"),
        ("Global index", "LME / resin market, daily. We forecast this", "largest mover for aluminium"),
        ("Premium over index", "Foil conversion, grade; contract, quarterly", "from ACG POs (pilot wk 3–4)"),
        ("Freight & insurance", "Container market, monthly; part of CIF", "counts toward the PVC floor"),
        ("USD/INR", "FX market, daily", "~1% of Al, ~36% of PVC swings"),
        ("Customs duty", "BCD 7.5% + AIDC (7.5% of BCD) + SWS 10%", "8.87% of CIF; 0% BCD if FTA"),
        ("Trade remedies", "Government decisions with dates", "foil ADD; PVC $766/t CIF floor")]
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        cell(7, 15, i, j, v)
for r in tbl.rows:
    r.height = Inches(0.29)
def mv(sid, top): shp(7, sid).top = Inches(top)
T(7, 16, "Port, clearance and inland haulage add a fixed ₹/t from ACG invoices. IGST is creditable,",
         "so it is cash flow, not cost. The duty rows are verified on ICEGATE before each PO.")
mv(16, 4.42)
T(7, 18, "How Smart Buy handles each gap")
T(7, 20, "Index: forecast six months ahead with an 80% band (the rest of this deck).")
T(7, 22, "Duty, FX and freight: calculated from current published rates and quotes.")
T(7, 24, "Premiums and port costs: fitted to ACG invoices, re-checked every quarter.")
T(7, 26, "Policy events: on a dated calendar, flagged in the buyer's note in advance.")
T(7, 27, "Three numbers that set the gap today")
T(7, 29, "8.87%"); T(7, 30, "customs duty on CIF value", "(MFN origin); IGST excluded")
T(7, 32, "$766/t"); T(7, 33, "CIF floor on suspension", "PVC, to late Jan 2027")
T(7, 35, "$479–721"); T(7, 36, "per t ADD on Chinese foil,", "extended to 15 Dec 2026")
T(7, 37, "What moves the landed price month to month")
T(7, 39, "Aluminium: the LME price itself. FX is only ~1% of monthly rupee-price variance.")
T(7, 41, "PVC: FX is ~36% of it, so the rupee view matters as much as the forecast.")
T(7, 43, "Origin decides duty: ASEAN, Japan, Korea pay 0% BCD; China can add ADD.")
T(7, 45, "Policy steps: PVC BCD was nil Apr–Jun 2026, then a $766/t floor from 24 Jul.")
T(7, 47, "Premiums reset quarterly with suppliers, so they are known in advance.")
T(7, 48, [("Worked example, PVC at the floor. ", True), ("CIF $766/t + BCD $57.4 + AIDC $4.3 + SWS $6.2", False)],
         "= $833.9/t duty-paid (+8.87%); at ₹95.45/$: ₹79.6/kg vs ₹73.1/kg CIF.",
         "Imports under Advance Authorisation for exports are exempt from the floor (conditions apply).")
T(7, 49, "Sources  Duty stack: CBIC/ICEGATE rates via EximPe (May 2026). PVC floor: DGFT Notification 25/2026-27, 24 Jul 2026. "
         "Foil ADD: Customs (ADD) extension to 15 Dec 2026, sunset review. FX: ECB. FX variance shares: this pipeline, 2010–2024.")
notes(7, "THE JUDGES' QUESTION: WHAT DRIVES THE DIFFERENCE VS THE GLOBAL INDEX?\n"
      "Answer first: five things sit between the index and ACG's price: premium, freight, FX, customs duty, and trade remedies. "
      "We forecast only the index. The other five are contracts, quotes, published rates or dated government decisions, so we calculate them.\n"
      "Then explain: for aluminium the index dominates (FX is about 1% of the monthly swing); for PVC, FX is about a third. "
      "Duty is 8.87% of CIF for a normal origin (BCD 7.5%, AIDC on BCD, SWS 10% on both), zero BCD from FTA origins. IGST is creditable, so not a cost. "
      "Two live events fall inside our pilot: the foil anti-dumping review (15 Dec 2026) and expiry of the PVC $766 floor (late Jan 2027). "
      "That makes the pilot a real test of the landed-cost model.")

# ============================ 8. PILOT PLAN (sharpened) ============================
T(8, 13, "90-day zero-capex shadow pilot, Nov 2026 – Feb 2027: what it costs, when it runs,",
         "and the numbers that decide go or stop")
T(8, 14, "Timeline, November 2026 to February 2027")
T(8, 17, "Weeks 1–2"); T(8, 18, "Data refresh and PO baseline")
T(8, 19, "Refresh public data; load 12–24 months of ACG POs for aluminium foil and PVC resin.")
T(8, 22, "Weeks 3–4"); T(8, 23, "Landed-cost overlay calibration")
T(8, 24, "Fit premiums, freight and port costs so the overlay matches ACG invoices; load the policy calendar.")
T(8, 27, "Months 2–3"); T(8, 28, "Shadow scoring, twice a month")
T(8, 29, "Buyers decide as today; every signal is logged against the actual buy. Spans the foil ADD and PVC floor decisions.")
T(8, 32, "Month 4"); T(8, 33, "Go / no-go review, Feb 2027")
T(8, 34, "Leadership decides against the gates on the right. If go: monthly run, then the ERP hook.")
T(8, 35, "Cost: ₹0 capex, no licences; runs on one laptop in about 20 minutes a month.",
         "ACG time: about 25 hours in total (sponsor 1 h/month, 2 buyers 30 min/fortnight, PO extract).")
T(8, 36, "Go gates at day 90 (all four must pass)")
gates = [("Gate", "Pass criterion"),
         ("Accuracy", "Forecast MAPE ≤ naïve at 1–3 months, both materials"),
         ("Landed cost", "Overlay within ±2.5% of ACG invoice on ≥ 80% of POs"),
         ("Decision value", "Shadow saving ≥ 0 after carry; full timing test at month 12"),
         ("Buyer trust", "≥ 80% of signals rated useful; every buy logged")]
for i, (a, b) in enumerate(gates):
    cell(8, 37, i, 0, a); cell(8, 37, i, 1, b)
T(8, 38, "Stop rules (automatic, any one triggers)")
stops = [("Trigger", "Automatic action"),
         ("Accuracy failure", "MAPE > naïve 2 months running → freeze signals, buy as today"),
         ("Landed-cost drift", "Overlay > 5% off invoices → pause and re-fit before resuming"),
         ("Regime break / stale data", "Band coverage < 50% or any feed > 45 days old → all BUY NORMAL")]
for i, (a, b) in enumerate(stops):
    cell(8, 39, i, 0, a); cell(8, 39, i, 1, b)
T(8, 40, "Why the savings verdict waits for month 12: the rule acts only a few times a quarter, so 90 days proves accuracy, "
         "landed-cost fit and adoption; proving savings needs about 12 months of signals.")
T(8, 41, "Governance: shadow mode throughout. Buyer decisions do not change during the pilot; nothing connects to ERP before the go decision.")
notes(8, "Answer first: zero capex, about 25 hours of ACG time, 90 days, four go gates and three automatic stop rules.\n"
      "If asked why not prove savings in 90 days: the rule acts a handful of times a quarter, so any 90-day saving would be luck either way. "
      "We prove what 90 days can prove: accuracy, landed-cost fit and adoption. The savings test runs to month 12.")

# ============================ 9. RISKS (bridge → policy calendar; ask row → takeaways) ============================
T(9, 13, "Risks we carry and how we contain them: policy events, residual model risk,", "and the roadmap after the pilot")
T(9, 14, "Policy events inside the pilot window")
T(9, 16, "15 Dec 2026 · Foil ADD (China, Malaysia, Thailand, Indonesia): sunset decision",
         "Late Jan 2027 · S-PVC $766/t CIF floor ends: extend, replace with ADD, or lapse",
         "Pending · S-PVC anti-dumping duty recommended by DGTR (Aug 2025)",
         "Pending · Countervailing-duty probe on Chinese S-PVC (opened Mar 2026)",
         "Past · PVC BCD nil 2 Apr – 30 Jun 2026, then restored",
         "Each shifts landed cost by a known amount on a known date,",
         "so the overlay calculates it; the forecast does not guess it.")
T(9, 17, "Why we forecast the index and calculate the rest")
T(9, 19, "Import parity: Indian producers price off import parity, so the global index also drives the domestic price ACG pays.")
T(9, 23, "Duties, floors and anti-dumping rates are published rules with dates: calculated exactly, not forecast.")
T(9, 42, "What leadership should take away")
T(9, 43, "1. Proven where it matters"); T(9, 44, "PVC timing beats random over 78 months; aluminium is reported honestly as near-efficient.")
T(9, 45, "2. Landed cost explained"); T(9, 46, "Index forecast plus calculated duty, FX, freight and premiums, checked against invoices.")
T(9, 47, "3. Low risk to try"); T(9, 48, "Zero capex, buyers keep deciding, and automatic stop rules end it at no cost.")
T(9, 49, "4. One decision needed"); T(9, 50, "Approve the shadow pilot from November 2026. Details on the next slide.")
T(9, 51, "Sources: [1] DGFT Notification 25/2026-27. [2] Customs (ADD) extension of foil duty to 15 Dec 2026. [3] DGTR final findings, S-PVC (Aug 2025). "
         "[4] DGTR CVD initiation, S-PVC (Mar 2026). [5] Audit trail: CHANGES.md in the repository.")

# ============================ 10. THE ASK (new closing slide) ============================
T(10, 6, "Our ask")
T(10, 7, "Approve a 90-day, zero-capex", "shadow pilot from November 2026")
T(10, 8, "Forecast the index. Calculate the rest.", "Let measured results, not judgement,", "decide when to buy ahead.")
T(10, 9, "Finals  ·  Technology Category")
T(10, 10, "Decision requested today")
T(10, 11, "T H E   T H R E E   T H I N G S   W E   N E E D")
T(10, 12, "One sponsor, one buyer per material, and purchase-order data. No capex, no licences, no systems change. "
          "A go / no-go review in February 2027 against gates we have set in advance.")
T(10, 14, "PO data"); T(10, 15, "12–24 months of aluminium foil", "and PVC resin purchase orders")
T(10, 17, "~25 hours"); T(10, 18, "of ACG time in total: a sponsor", "and one buyer per material")
T(10, 20, "Go"); T(10, 21, "approval to start the shadow", "pilot in November 2026")
T(10, 23, "W H A T   A C G   G E T S   B Y   F E B R U A R Y   2 0 2 7")
T(10, 24, "A measured answer to three questions. Does the forecast beat a no-change baseline on ACG's own buys? Does the "
          "landed-cost model match ACG invoices through two live policy decisions? Do buyers trust and use the signal? "
          "If yes, Smart Buy runs monthly and connects to the ERP. If not, the stop rules end it at no cost.")
T(10, 28, "Pilot design: slide 8. Landed-cost model: slide 7. Technical appendix: A1–A7. Q&A welcome.")
cell(10, 27, 4, 1, "Live dashboard, 14-slide deck + appendix, reproducible pipeline with leakage tests")
notes(10, "CLOSE, 30 SECONDS. Do not end on 'thank you'; end on the decision.\n"
      "We are asking for three things: purchase-order data, about 25 hours of your team's time, and a go to start in November. "
      "In February you get a measured answer: does it beat a no-change forecast on your buys, does the landed-cost model match your invoices "
      "through two live policy decisions, and do your buyers use it. If not, it stops automatically at no cost. "
      "Can we have your approval to start the pilot in November?")

# ============================ NAV BAR: match the new structure ============================
MAIN_TABS = ["Problem", "Solution", "Proof", "Demo", "Landed cost", "Pilot", "Risks & ask"]
MAIN_ACTIVE = {2: 0, 3: 1, 4: 2, 5: 2, 6: 3, 7: 4, 8: 5, 9: 6}
APPX_TABS = ["A1 Value chain", "A2 Data & lags", "A3 Drivers", "A4 Models", "A5 Explainability", "A6 Validation", "A7 Bands"]
def set_nav(n, labels, active):
    sl = S[n]
    pills = sorted([sh for sh in sl.shapes if sh.has_text_frame and sh.top < Inches(0.6) and sh.width < Inches(1.9)
                    and sh.text_frame.text.strip()], key=lambda sh: sh.left)
    assert len(pills) == 7, (n, len(pills))
    red = [sh for sh in sl.shapes if sh.top < Inches(0.5) and sh.shape_type == 1 and sh.fill.type == 1
           and str(sh.fill.fore_color.rgb) == "C8102E"][0]
    off = Inches(0.03)
    red.left = pills[active].left + off
    first = min(pills, key=lambda sh: list(sl.shapes._spTree).index(sh._element))
    first._element.addprevious(red._element)          # highlight sits behind every label
    for i, (sh, lab) in enumerate(zip(pills, labels)):
        _fill_body(sh.text_frame._txBody, [lab], "FFFFFF" if i == active else "4A4D55")
for n, a in MAIN_ACTIVE.items():
    set_nav(n, MAIN_TABS, a)
for n in range(11, 18):
    set_nav(n, APPX_TABS, n - 11)

# ============================ FINAL ORDER: the evidence is in the main flow ============================
# business titles + presenter notes for the four slides that move from the appendix into the main story
T(13, 13, "What drives the price: aluminium moves with copper and energy in the same month;",
          "crude oil and China demand give PVC one to two months' warning")
notes(13, "Answer first: for PVC, crude oil and China's leading indicator move one to two months ahead of the price, so they give buyers warning. "
          "For aluminium, copper and energy move in the same month: they explain the price but cannot forecast it, which is why aluminium is hard to beat.")
notes(16, "Answer first: we tested it the way a buyer would live with it. Every month from 2018, the model was retrained on only what was public that day, "
          "then forecast the next six months. The red line is the forecast made three months before the black line happened. "
          "Skill is highest at one to two months, which is where most buying decisions sit.")
notes(14, "Answer first: ten models compete on the same test; a model gets a vote only in months when it has beaten the no-change forecast on its own past record. "
          "The best single model is N-BEATS, a neural network trained on 74 commodities at once, because one price alone has too little history.")
notes(15, "Answer first: every forecast comes with a short note a buyer can read in ten seconds: the expected move, the range, the chance of a rise, "
          "the signal, and the three drivers that pushed it most.")

FINAL = [1, 2, 3, 13, 16, 4, 14, 5, 6, 15, 7, 8, 9, 10, 11, 12, 17]
_lst = prs.slides._sldIdLst
_ids = list(_lst)
for _e in _ids:
    _lst.remove(_e)
for _n in FINAL:
    _lst.append(_ids[_n - 1])
S = {i + 1: s_ for i, s_ in enumerate(prs.slides)}
NMAIN = 14
SEC2 = {2: "P R O B L E M", 3: "S O L U T I O N", 4: "W H A T   D R I V E S   T H E   P R I C E", 5: "H O W   W E   T E S T E D",
        6: "R E S U L T S   V S   G A T E S", 7: "W H I C H   M O D E L   W I N S", 8: "P R O O F   I N   R U P E E S",
        9: "W O R K I N G   D E M O", 10: "W H Y   T H E   F O R E C A S T   M O V E D", 11: "L A N D E D   C O S T",
        12: "P I L O T   P L A N", 13: "R I S K S   &   R O A D M A P"}
APX2 = {15: "V A L U E   C H A I N", 16: "D A T A   &   L E A K A G E   C O N T R O L S", 17: "U N C E R T A I N T Y   B A N D S"}
for n, sl in S.items():
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if re.fullmatch(r"\d+ / \d+|A\d+", t) and sh.top > Inches(6.8):   # footer page number only
            T(n, sh.shape_id, f"{n} / {NMAIN}" if n <= NMAIN else f"A{n - NMAIN}")
        if sh.shape_id == 12 and "·" in t and n not in (1, NMAIN):
            T(n, 12, f"{' '.join(f'{n:02d}')}  ·  {SEC2[n]}" if n < NMAIN else f"A {n - NMAIN}  ·  A P P E N D I X  ·  {APX2[n]}")
for n, a in {2: 0, 3: 1, 4: 1, 5: 2, 6: 2, 7: 2, 8: 2, 9: 3, 10: 3, 11: 4, 12: 5, 13: 6}.items():
    set_nav(n, MAIN_TABS, a)
for n in (15, 16, 17):
    set_nav(n, ["A1 Value chain", "A2 Data & leakage", "A3 Bands", " ", " ", " ", " "], n - 15)
T(14, 28, "Pilot design: slide 12. Landed-cost model: slide 11. Backup: appendix A1–A3. Q&A welcome.")

# ============================ LIVE NUMBERS: slide text always matches the outputs ============================
import importlib.util as _ilu
_spec = _ilu.spec_from_file_location("live_numbers", str(Path(DECK_ROOT) / "src" / "deck" / "live_numbers.py"))
_ln = _ilu.module_from_spec(_spec); _spec.loader.exec_module(_ln)
LIVE = {k: v for k, v in _ln.compute().items() if k != v}
def _swap(t):
    for old in sorted(LIVE, key=len, reverse=True):
        t = t.replace(old, LIVE[old])
    return t
for sl in prs.slides:
    for r in sl._element.iter(qn("a:t")):
        if r.text: r.text = _swap(r.text)
    if sl.has_notes_slide:
        for r in sl.notes_slide._element.iter(qn("a:t")):
            if r.text: r.text = _swap(r.text)
print("live numbers changed:", len(LIVE))
prs.save(str(SB / "outputs" / "SmartBuy_Finals_Deck.pptx"))
print("saved")
