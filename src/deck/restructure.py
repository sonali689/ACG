"""Step between the 15-slide v2 deck and the finals deck:
duplicate slide 4 (becomes the landed-cost slide) and slide 1 (becomes the closing ask slide),
then reorder into the leadership flow (10 main slides + 7 appendix slides).
Plain python-pptx; neither source slide contains pictures, so a deep copy of its shapes is complete."""
from copy import deepcopy
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn

SB = Path(__file__).resolve().parents[2]
SRC = SB / "outputs" / "deck" / "SmartBuy_Round2_15Slides_v2.pptx"
DST = SB / "outputs" / "deck" / "_restructured.pptx"

prs = Presentation(str(SRC))


def duplicate(index):
    src = prs.slides[index]
    new = prs.slides.add_slide(src.slide_layout)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is not None:
        new._element.cSld.insert(0, deepcopy(bg))
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        new.shapes._spTree.append(deepcopy(el))
    return new


duplicate(3)   # -> position 16: landed-cost slide (from the data/lag slide layout)
duplicate(0)   # -> position 17: closing ask slide (from the title slide layout)

order = [1, 2, 6, 10, 12, 13, 16, 14, 15, 17, 3, 4, 5, 7, 8, 9, 11]   # 1-based, old positions
lst = prs.slides._sldIdLst
ids = list(lst)
for el in ids:
    lst.remove(el)
for n in order:
    lst.append(ids[n - 1])
prs.save(str(DST))
print("restructured ->", DST.name)
